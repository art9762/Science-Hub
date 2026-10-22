"""Bounded document extraction and portable workspace exports; no embedded execution."""

from __future__ import annotations

import io
import mimetypes
import re
import shutil
import warnings
import zipfile
from html import escape
from html.parser import HTMLParser
from pathlib import Path

MAX_UPLOAD_BYTES = 30 * 1024 * 1024
MAX_EXTRACTED_CHARS = 3_000_000
MAX_PAGES = 500
MAX_OCR_PAGES = 12
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".tif", ".tiff", ".bmp"}
TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".json", ".log", ".rst", ".yaml", ".yml", ".xml"}
SUPPORTED_EXTENSIONS = IMAGE_EXTENSIONS | TEXT_EXTENSIONS | {".pdf", ".docx", ".html", ".htm"}


def safe_filename(name: str) -> str:
    name = str(name or "document").replace("\\", "/").split("/")[-1]
    name = re.sub(r'[\x00-\x1f\x7f<>:"|?*]', "_", name).strip(" .")
    return (name or "document")[:200]


class PlainHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.hidden = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript", "template"}:
            self.hidden += 1
        elif tag in {"p", "br", "div", "li", "h1", "h2", "h3", "tr", "section"} and not self.hidden:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript", "template"}:
            self.hidden = max(0, self.hidden - 1)
        elif tag in {"p", "div", "li", "h1", "h2", "h3", "tr", "section"} and not self.hidden:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1251", errors="replace")


def _ocr(image) -> tuple[str, str]:
    if not shutil.which("tesseract"):
        return "", "unavailable"
    try:
        import pytesseract

        available = pytesseract.get_languages(config="")
        language = "+".join(item for item in ("rus", "eng") if item in available) or (
            available[0] if available else "eng"
        )
        return pytesseract.image_to_string(image, lang=language, timeout=20).strip(), "completed"
    except ImportError:
        return "", "unavailable"
    except Exception:
        return "", "failed"


def _pdf(path: Path) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("PDF защищён паролем. Загрузите незашифрованную копию.")
    if len(reader.pages) > MAX_PAGES:
        raise ValueError(f"Допустимо не более {MAX_PAGES} страниц в PDF.")
    pages, total = [], 0
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        total += len(text)
        if total > MAX_EXTRACTED_CHARS:
            raise ValueError("Слишком много текста в документе: лимит 3 млн символов.")
        pages.append({"number": number, "text": text})
    empty = [page for page in pages if not page["text"].strip()]
    ocr_status = "not_needed"
    notices = []
    if empty:
        ocr_status = "unavailable"
        try:
            import fitz
            from PIL import Image

            if shutil.which("tesseract"):
                with fitz.open(path) as rendered:
                    for page in empty[:MAX_OCR_PAGES]:
                        source_page = rendered[page["number"] - 1]
                        if source_page.rect.width * source_page.rect.height * 2.25 > 25_000_000:
                            notices.append(f"Страница {page['number']}: превышен предел размера OCR.")
                            ocr_status = "partial"
                            continue
                        pixmap = source_page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                        image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                        page["text"], ocr_status = _ocr(image)
                        page["extraction"] = "ocr"
                if len(empty) > MAX_OCR_PAGES:
                    ocr_status = "partial"
                    notices.append(f"OCR ограничен {MAX_OCR_PAGES} страницами за загрузку.")
        except ImportError:
            pass
        except Exception:
            ocr_status = "failed"
        if ocr_status == "unavailable":
            notices.append("Для сканированных страниц установите PyMuPDF, pytesseract и Tesseract OCR.")
    return {
        "pages": pages,
        "ocr_status": ocr_status,
        "warnings": notices,
        "metadata": {"page_count": len(pages)},
        "status": "ready" if any(page["text"].strip() for page in pages) else "needs_ocr",
    }


def _docx(path: Path) -> dict:
    from docx import Document

    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10_000 or sum(info.file_size for info in entries) > 100 * 1024 * 1024:
            raise ValueError("DOCX превышает лимит распакованного размера.")
        if any(
            info.file_size > 1_000_000 and info.file_size > max(info.compress_size, 1) * 200
            for info in entries
        ):
            raise ValueError("Недопустимая степень сжатия DOCX.")
    document = Document(path)
    parts = []
    # Preserve paragraph/table reading order. DOCX has no stable page map until rendered.
    for element in document.iter_inner_content():
        if hasattr(element, "text"):
            parts.append(element.text)
        elif hasattr(element, "rows"):
            parts.extend(" | ".join(cell.text for cell in row.cells) for row in element.rows)
    return {
        "pages": [{"number": 1, "text": "\n".join(parts)}],
        "metadata": {"page_count": 1, "page_layout": "logical"},
        "warnings": ["DOCX показан единым текстом: физические страницы зависят от редактора."],
        "status": "ready",
        "ocr_status": "not_needed",
    }


def _image(path: Path) -> dict:
    from PIL import Image, ImageOps

    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(path) as original:
            width, height = original.size
            if width * height > 40_000_000:
                raise ValueError("Изображение превышает лимит 40 мегапикселей.")
            metadata = {
                "width": width,
                "height": height,
                "format": original.format,
                "frames": getattr(original, "n_frames", 1),
            }
            mime_type = Image.MIME.get(
                original.format, mimetypes.guess_type(path.name)[0] or "application/octet-stream"
            )
            original.seek(0)
            image = ImageOps.exif_transpose(original).convert("RGB")
            image.thumbnail((3500, 3500))
            text, status = _ocr(image)
    notices = []
    if status == "unavailable":
        notices.append(
            "OCR недоступен. Для поиска текста установите Tesseract и pytesseract; описание изображения создаётся отдельно моделью зрения."
        )
    elif status == "failed":
        notices.append("OCR не завершён. Изображение сохранено и доступно для анализа моделью зрения.")
    return {
        "pages": [{"number": 1, "text": text, "extraction": "ocr"}],
        "metadata": metadata,
        "ocr_status": status,
        "status": "ready",
        "warnings": notices,
        "mime_type": mime_type,
    }


def parse_document(path: Path, name: str) -> dict:
    extension = Path(name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Формат не поддерживается. Загрузите PDF, DOCX, текст, Markdown, HTML или изображение."
        )
    try:
        if extension == ".pdf":
            result = _pdf(path)
        elif extension == ".docx":
            result = _docx(path)
        elif extension in IMAGE_EXTENSIONS:
            result = _image(path)
        else:
            text = _decode(path.read_bytes())
            if "\x00" in text:
                raise ValueError("Файл не похож на текстовый документ.")
            if extension in {".html", ".htm"}:
                parser = PlainHTML()
                parser.feed(text)
                text = re.sub(r"\n{3,}", "\n\n", "".join(parser.parts)).strip()
            result = {
                "pages": [{"number": 1, "text": text}],
                "status": "ready",
                "metadata": {"page_count": 1},
                "ocr_status": "not_needed",
                "warnings": [],
            }
        result["text"] = "\n\n".join(page["text"] for page in result["pages"])
        if len(result["text"]) > MAX_EXTRACTED_CHARS:
            raise ValueError("Слишком много текста в документе: лимит 3 млн символов.")
        result.setdefault("mime_type", mimetypes.guess_type(name)[0] or "application/octet-stream")
        result["description"] = ""
        return result
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Не удалось прочитать документ. Проверьте формат и целостность файла.") from exc


def render_pdf_page(path: Path, page_number: int) -> bytes:
    """Rasterize original page content without plugins or embedded script execution."""
    import math
    import fitz

    with fitz.open(path) as document:
        if not document.is_pdf:
            raise ValueError("Просмотр страниц доступен только для PDF.")
        if page_number < 1 or page_number > len(document):
            raise IndexError("Страница отсутствует в документе.")
        page = document[page_number - 1]
        width, height = page.rect.width, page.rect.height
        if not all(math.isfinite(value) and value > 0 for value in (width, height)):
            raise ValueError("Некорректный размер страницы PDF.")
        scale = min(2.0, 1600 / max(width, height))
        pixels = page.get_pixmap(matrix=fitz.Matrix(scale, scale), colorspace=fitz.csRGB, alpha=False)
        return pixels.tobytes("png")


def _plain_markdown(content: str) -> str:
    text = re.sub(r"(?m)^#{1,6}\s+", "", content)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    return text


def _unicode_pdf_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    name = "ScienceHubUnicode"
    if name in pdfmetrics.getRegisteredFontNames():
        return name
    paths = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    for path in paths:
        if path.is_file():
            try:
                pdfmetrics.registerFont(TTFont(name, str(path)))
                return name
            except Exception:
                continue
    raise ValueError("Для экспорта PDF нужен Unicode-шрифт: установите DejaVu Sans или Arial.")


def export_document(document: dict, format: str) -> tuple[bytes, str, str]:
    title, content = document.get("title", "Документ"), document.get("content", "")
    if format == "md":
        return content.encode("utf-8"), "text/markdown; charset=utf-8", ".md"
    if format == "txt":
        return _plain_markdown(content).encode("utf-8"), "text/plain; charset=utf-8", ".txt"
    output = io.BytesIO()
    if format == "docx":
        from docx import Document

        exported = Document()
        exported.core_properties.title = title
        for line in content.splitlines():
            heading = re.match(r"^(#{1,6})\s+(.*)$", line)
            if heading:
                exported.add_heading(heading[2], len(heading[1]))
            elif line.startswith(("- ", "* ")):
                exported.add_paragraph(line[2:], style="List Bullet")
            else:
                exported.add_paragraph(line)
        exported.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".docx",
        )
    if format == "pdf":
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        font = _unicode_pdf_font()
        body = ParagraphStyle(
            "Body", fontName=font, fontSize=10, leading=15, alignment=TA_LEFT, splitLongWords=True
        )
        heading_style = ParagraphStyle("Heading", parent=body, fontSize=16, leading=21, spaceAfter=9)
        story = []
        for line in content.splitlines():
            heading = re.match(r"^#{1,6}\s+(.*)$", line)
            if not line.strip():
                story.append(Spacer(1, 8))
            else:
                story.append(
                    Paragraph(escape(heading[1] if heading else line), heading_style if heading else body)
                )
        if not story:
            story.append(Paragraph(escape(title), heading_style))
        SimpleDocTemplate(
            output,
            pagesize=A4,
            title=title,
            author="Science Hub",
            leftMargin=48,
            rightMargin=48,
            topMargin=48,
            bottomMargin=48,
        ).build(story)
        return output.getvalue(), "application/pdf", ".pdf"
    raise ValueError("Формат экспорта должен быть md, txt, docx или pdf.")

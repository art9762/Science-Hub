from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.responses import JSONResponse
from fastapi.testclient import TestClient

from backend.core import router
from backend.documents import export_document, parse_document
from backend.storage import Store


@pytest.fixture
def client(tmp_path):
    app = FastAPI()
    app.state.store = Store(tmp_path)
    app.include_router(router, prefix="/api")

    @app.exception_handler(KeyError)
    async def not_found(request, exc):
        return JSONResponse({"error": "not found"}, status_code=404)

    with TestClient(app) as client:
        yield client


def project(client, name="Проект"):
    response = client.post("/api/projects", json={"name": name})
    assert response.status_code == 201
    return response.json()["id"]


def upload(client, pid, text="Первый факт. Второй факт.", filename="knowledge.txt"):
    response = client.post(
        f"/api/projects/{pid}/documents", files={"files": (filename, text.encode("utf-8"))}
    )
    assert response.status_code == 201, response.text
    return response.json()[0]


def test_store_durable_isolation_and_concurrent_writes(tmp_path):
    store = Store(tmp_path)
    first = store.create_project("Один")
    second = store.create_project("Два")
    document = store.put(first["id"], "document", {"name": "secret.txt", "text": "secret"})
    with pytest.raises(KeyError):
        store.get(second["id"], "document", document["id"])
    with pytest.raises(KeyError):
        store.put(second["id"], "document", {**document, "project_id": second["id"]})
    with pytest.raises(KeyError):
        store.put(first["id"], "claim", {"id": document["id"], "text": "wrong type"})
    with ThreadPoolExecutor(max_workers=4) as executor:
        records = list(
            executor.map(lambda index: store.put(first["id"], "event", {"message": str(index)}), range(20))
        )
    assert len({record["id"] for record in records}) == 20
    reopened = Store(tmp_path)
    assert reopened.get(first["id"], "document", document["id"])["text"] == "secret"
    assert len(reopened.list(first["id"], "event")) == 20
    assert reopened.list(second["id"], "event") == []
    replaced = reopened.put(first["id"], "document", {"id": document["id"], "name": "new.txt"})
    assert "text" not in replaced
    assert replaced["created_at"] == document["created_at"]
    reopened.delete_project(first["id"])
    assert reopened.project(second["id"])["name"] == "Два"


def test_project_seed_empty_state_and_names(client):
    pid = project(client)
    state = client.get(f"/api/projects/{pid}/state").json()
    assert len(state["agents"]) == 8
    assert {agent["role"] for agent in state["agents"]} == {
        "engineer",
        "analyst",
        "researcher",
        "critic",
        "fact_checker",
        "technologist",
        "estimator",
        "pessimist",
    }
    assert state["documents"] == state["claims"] == state["messages"] == []
    assert client.post("/api/projects", json={"name": "  "}).status_code == 422
    assert client.patch(f"/api/projects/{pid}", json={"name": "Название"}).json()["name"] == "Название"


def test_annotations_preserve_page_offsets_and_search_all_sources(client):
    pid = project(client)
    document = upload(client, pid)
    assert "file_path" not in document and "text" not in document and "pages" not in document
    did = document["id"]
    full = client.get(f"/api/projects/{pid}/documents/{did}").json()
    assert full["pages"][0]["text"].startswith("Первый факт")
    assert "file_path" not in full
    annotation = {
        "document_id": did,
        "page": 1,
        "start": 0,
        "end": 6,
        "quote": "Первый",
        "note": "Проверить источник",
        "type": "verify",
    }
    saved = client.post(f"/api/projects/{pid}/annotations", json=annotation)
    assert saved.status_code == 201, saved.text
    assert (
        client.post(f"/api/projects/{pid}/annotations", json={**annotation, "quote": "Второй"}).status_code
        == 422
    )
    assert client.post(f"/api/projects/{pid}/annotations", json={**annotation, "page": 2}).status_code == 422
    assert client.post(f"/api/projects/{pid}/annotations", json={**annotation, "end": 999}).status_code == 422
    client.patch(
        f"/api/projects/{pid}/documents/{did}",
        json={"description": "Схема насосной станции", "category": "Энергетика"},
    )
    assert client.get(f"/api/projects/{pid}/search", params={"q": "ВТОРОЙ"}).json()[0]["page"] == 1
    annotation_hit = client.get(f"/api/projects/{pid}/search", params={"q": "источник"}).json()[0]
    assert annotation_hit["annotation_id"] == saved.json()["id"]
    assert annotation_hit["start"] == 0 and annotation_hit["end"] == 6
    assert annotation_hit["quote"] == "Первый"
    assert (
        client.get(f"/api/projects/{pid}/search", params={"q": "насосной"}).json()[0]["source"]
        == "description"
    )
    assert (
        client.get(f"/api/projects/{pid}/search", params={"q": "Энергетика"}).json()[0]["source"]
        == "metadata"
    )


def test_cross_project_references_and_cascade(client):
    one, two = project(client, "Один"), project(client, "Два")
    document = upload(client, one)
    did = document["id"]
    assert client.get(f"/api/projects/{two}/documents/{did}").status_code == 404
    assert client.delete(f"/api/projects/{two}/documents/{did}").status_code == 404
    assert (
        client.get(f"/api/projects/{two}/search", params={"q": "факт", "document_ids": did}).status_code
        == 404
    )
    assert (
        client.post(
            f"/api/projects/{two}/annotations", json={"document_id": did, "note": "other"}
        ).status_code
        == 404
    )
    assert (
        client.post(f"/api/projects/{two}/claims", json={"text": "Факт", "source_ids": [did]}).status_code
        == 404
    )
    assert (
        client.post(
            f"/api/projects/{two}/board-items", json={"type": "document", "reference_id": did}
        ).status_code
        == 404
    )
    annotation = client.post(f"/api/projects/{one}/annotations", json={"document_id": did, "note": "Заметка"})
    assert annotation.status_code == 201
    assert (
        client.post(
            f"/api/projects/{one}/board-items", json={"type": "document", "reference_id": did}
        ).status_code
        == 201
    )
    claim = client.post(f"/api/projects/{one}/claims", json={"text": "Факт", "source_ids": [did]}).json()
    assert claim["confidence"] is None and claim["confidence_provenance"] == "user_assessment"
    link = client.post(f"/api/projects/{one}/links", json={"source": did, "target": claim["id"]})
    assert link.status_code == 201
    assert client.delete(f"/api/projects/{one}/documents/{did}").status_code == 204
    state = client.get(f"/api/projects/{one}/state").json()
    assert state["annotations"] == state["board_items"] == state["links"] == []
    assert state["claims"][0]["status"] == "unverified"
    assert state["claims"][0]["confidence"] is None
    assert state["claims"][0]["source_ids"] == []
    assert state["claims"][0]["missing_source_ids"] == [did]
    assert "Источник удалён" in state["claims"][0]["reason"]
    assert not any((client.app.state.store.root / "files" / one).iterdir())


def test_upload_traversal_html_safety_and_batch_rollback(client):
    pid = project(client)
    html = "<h1>Название</h1><p>Обычный текст</p><script>window.steal()</script><style>hide</style>"
    doc = upload(client, pid, html, "../../hostile.html")
    assert doc["name"] == "hostile.html"
    full = client.get(f"/api/projects/{pid}/documents/{doc['id']}").json()
    assert "Обычный текст" in full["text"] and "steal" not in full["text"]
    response = client.get(f"/api/projects/{pid}/documents/{doc['id']}/file")
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    assert response.headers["content-security-policy"] == "sandbox"
    response = client.post(
        f"/api/projects/{pid}/documents",
        files=[("files", ("valid.txt", b"good")), ("files", ("bad.exe", b"bad"))],
    )
    assert response.status_code == 422
    assert len(client.get(f"/api/projects/{pid}/state").json()["documents"]) == 1
    with pytest.raises(ValueError):
        client.app.state.store.file_path("../outside.txt")
    with pytest.raises(ValueError):
        client.app.state.store.file_path("/etc/passwd")


def test_decisions_cycles_and_partial_updates(client):
    pid = project(client)
    parent = client.post(f"/api/projects/{pid}/decisions", json={"title": "Материал"}).json()
    child = client.post(
        f"/api/projects/{pid}/decisions", json={"title": "Сталь", "parent_id": parent["id"]}
    ).json()
    assert (
        client.patch(
            f"/api/projects/{pid}/decisions/{parent['id']}", json={"parent_id": child["id"]}
        ).status_code
        == 422
    )
    accepted = client.patch(
        f"/api/projects/{pid}/decisions/{child['id']}", json={"status": "accepted"}
    ).json()
    assert accepted["title"] == "Сталь" and accepted["parent_id"] == parent["id"]
    assert (
        client.patch(f"/api/projects/{pid}/decisions/{child['id']}", json={"id": "forged"}).status_code == 422
    )
    assert (
        client.post(f"/api/projects/{pid}/claims", json={"text": "Факт", "confidence": 90}).status_code == 422
    )
    client.delete(f"/api/projects/{pid}/decisions/{parent['id']}")
    assert client.get(f"/api/projects/{pid}/decisions").json()[0]["parent_id"] is None


def test_workspace_exports_roundtrip_unicode_and_scoping(client):
    from docx import Document
    from pypdf import PdfReader

    one, two = project(client), project(client)
    content = "# Исследование\nНасос и двигатель: 42 кВт.\n\n- Проверка температуры"
    result = client.post(
        f"/api/projects/{one}/workspace-documents", json={"title": "Отчёт", "content": content}
    )
    assert result.status_code == 201
    wid = result.json()["id"]
    for extension in ("md", "txt", "docx", "pdf"):
        response = client.get(
            f"/api/projects/{one}/workspace-documents/{wid}/export", params={"format": extension}
        )
        assert response.status_code == 200, response.text
        assert "filename*=" in response.headers["content-disposition"]
        if extension == "md":
            assert response.text == content
        elif extension == "txt":
            assert response.text.startswith("Исследование")
        elif extension == "docx":
            assert "Насос" in "\n".join(p.text for p in Document(BytesIO(response.content)).paragraphs)
        else:
            assert "Насос" in PdfReader(BytesIO(response.content)).pages[0].extract_text()
    assert client.get(f"/api/projects/{two}/workspace-documents/{wid}/export").status_code == 404


def test_real_pdf_docx_image_extraction_and_ocr_unavailable(tmp_path, monkeypatch):
    from docx import Document
    from PIL import Image

    pdf, _, _ = export_document({"title": "PDF", "content": "# Проверка\nПрочность конструкции"}, "pdf")
    path = tmp_path / "test.pdf"
    path.write_bytes(pdf)
    assert "Прочность" in parse_document(path, "test.pdf")["text"]
    word = Document()
    word.add_paragraph("До таблицы")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Материал"
    table.cell(0, 1).text = "Сталь"
    word.add_paragraph("После таблицы")
    path = tmp_path / "test.docx"
    word.save(path)
    result = parse_document(path, path.name)
    assert result["text"] == "До таблицы\nМатериал | Сталь\nПосле таблицы"
    path = tmp_path / "image.png"
    Image.new("RGB", (32, 20), "white").save(path)
    monkeypatch.setattr("backend.documents.shutil.which", lambda command: None)
    result = parse_document(path, path.name)
    assert result["metadata"]["width"] == 32
    assert result["description"] == "" and result["ocr_status"] == "unavailable"
    assert result["warnings"]


def test_page_two_anchors_and_project_file_cleanup(client):
    from reportlab.pdfgen import canvas

    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(40, 700, "First page")
    pdf.showPage()
    pdf.drawString(40, 700, "Second page bearing")
    pdf.save()
    pid = project(client)
    response = client.post(f"/api/projects/{pid}/documents", files={"files": ("two.pdf", output.getvalue())})
    assert response.status_code == 201
    did = response.json()[0]["id"]
    hit = client.get(f"/api/projects/{pid}/search", params={"q": "bearing"}).json()[0]
    assert hit["page"] == 2
    assert (
        client.post(
            f"/api/projects/{pid}/annotations",
            json={
                "document_id": did,
                "page": 2,
                "start": 0,
                "end": 6,
                "quote": "Second",
                "type": "important",
            },
        ).status_code
        == 201
    )
    directory = client.app.state.store.root / "files" / pid
    assert directory.exists()
    client.delete(f"/api/projects/{pid}")
    assert not directory.exists()


def test_real_ocr_image_and_scanned_pdf(tmp_path):
    import shutil
    from pathlib import Path
    from PIL import Image, ImageDraw, ImageFont

    if not shutil.which("tesseract"):
        pytest.skip("Optional Tesseract not installed")
    pytest.importorskip("pytesseract")
    pytest.importorskip("fitz")
    fonts = [
        Path("/Library/Fonts/Arial Unicode.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    font_path = next((path for path in fonts if path.exists()), None)
    if not font_path:
        pytest.skip("No test font available")
    image = Image.new("RGB", (1200, 220), "white")
    ImageDraw.Draw(image).text(
        (40, 70), "ENGINEERING 12345", fill="black", font=ImageFont.truetype(str(font_path), 64)
    )
    path = tmp_path / "ocr.png"
    image.save(path)
    result = parse_document(path, path.name)
    assert result["ocr_status"] == "completed"
    assert "12345" in result["text"] and "ENGINEERING" in result["text"].upper()
    from reportlab.pdfgen import canvas

    pdf_path = tmp_path / "scan.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=(600, 220))
    pdf.drawImage(str(path), 0, 40, width=600, height=110)
    pdf.save()
    result = parse_document(pdf_path, pdf_path.name)
    assert result["status"] == "ready"
    assert result["ocr_status"] == "completed"
    assert "12345" in result["pages"][0]["text"]


def test_project_deletion_cancels_its_background_tasks(client):
    import asyncio

    pid = project(client)
    store = client.app.state.store
    run = store.put(pid, "run", {"status": "running"})
    completed = []

    async def background():
        try:
            await asyncio.Event().wait()
        finally:
            completed.append("cancelled")

    async def start():
        client.app.state.ai_tasks = {run["id"]: asyncio.create_task(background())}

    client.portal.call(start)
    assert client.delete(f"/api/projects/{pid}").status_code == 204
    assert completed == ["cancelled"]


def test_image_upload_schedules_analysis_after_persistence(client, monkeypatch):
    from PIL import Image

    pid = project(client)
    scheduled = []

    def schedule(app, project_id, document_id):
        document = app.state.store.get(project_id, "document", document_id)
        assert app.state.store.file_path(document["file_path"]).is_file()
        app.state.store.put(project_id, "document", {**document, "vision_status": "pending_settings"})
        scheduled.append(document_id)

    monkeypatch.setattr("backend.ai.schedule_image_analysis", schedule)
    monkeypatch.setattr("backend.documents.shutil.which", lambda command: None)
    output = BytesIO()
    Image.new("RGB", (16, 16), "white").save(output, format="PNG")
    response = client.post(
        f"/api/projects/{pid}/documents", files={"files": ("image.png", output.getvalue())}
    )
    assert response.status_code == 201
    assert scheduled == [response.json()[0]["id"]]
    assert response.json()[0]["vision_status"] == "pending_settings"


def test_search_unicode_anchors_and_nested_cascade(client):
    pid = project(client)
    doc = upload(client, pid, "Straße ENGINEERING Straße")
    result = client.get(f"/api/projects/{pid}/search", params={"q": "engineering"}).json()[0]
    assert result["start"] == 7 and result["end"] == 18 and result["quote"] == "ENGINEERING"
    store = client.app.state.store
    card = store.put(pid, "board_item", {"reference_id": doc["id"]})
    claim = store.put(
        pid, "claim", {"text": "claim", "status": "supported", "confidence": 0.9, "source_ids": [doc["id"]]}
    )
    store.put(pid, "link", {"source": card["id"], "target": claim["id"]})
    store.delete(pid, "document", doc["id"])
    assert store.list(pid, "link") == []
    claim = store.get(pid, "claim", claim["id"])
    assert claim["status"] == "unverified" and claim["confidence"] is None


def test_original_pdf_page_raster_is_real_bounded_and_project_scoped(client):
    from PIL import Image
    from reportlab.pdfgen import canvas

    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=(2000, 1000))
    pdf.setFillColorRGB(1, 0, 0)
    pdf.rect(0, 0, 2000, 1000, fill=1, stroke=0)
    pdf.showPage()
    pdf.setFillColorRGB(0, 0, 1)
    pdf.rect(0, 0, 2000, 1000, fill=1, stroke=0)
    pdf.save()
    one, two = project(client), project(client)
    response = client.post(
        f"/api/projects/{one}/documents", files={"files": ("original.pdf", output.getvalue())}
    )
    assert response.status_code == 201
    did = response.json()[0]["id"]
    route = f"/api/projects/{one}/documents/{did}/pages"
    first = client.get(f"{route}/1/image")
    assert first.status_code == 200
    assert first.headers["content-type"] == "image/png"
    assert first.content.startswith(b"\x89PNG\r\n\x1a\n")
    with Image.open(BytesIO(first.content)) as image:
        assert max(image.size) <= 1600
        assert image.getpixel((image.width // 2, image.height // 2)) == (255, 0, 0)
    second = client.get(f"{route}/2/image")
    with Image.open(BytesIO(second.content)) as image:
        assert image.getpixel((image.width // 2, image.height // 2)) == (0, 0, 255)
    assert client.get(f"{route}/0/image").status_code == 404
    assert client.get(f"{route}/3/image").status_code == 404
    assert client.get(f"/api/projects/{two}/documents/{did}/pages/1/image").status_code == 404
    text = upload(client, one)
    assert client.get(f"/api/projects/{one}/documents/{text['id']}/pages/1/image").status_code == 422
